from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "report_work"
OUTPUT = WORK / "RM75_Stage1_强化学习汇报总结_优化版.docx"
FONT_SANS = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_SERIF = "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"

NAVY = "1F4E79"
BLUE = "DCE6F1"
PALE = "F3F6FA"
GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    """Keep each table row on one page so a record is not broken in half."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def remove_paragraph_border(paragraph) -> None:
    """Remove borders inherited from a paragraph style, especially on the cover."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_run_font(run, name="Noto Sans CJK SC", size=10.5, bold=False, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=5, line=1.25, keep=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    fmt.widow_control = True


def add_text(doc, text: str, *, bold_lead: str | None = None, style=None, after=5) -> None:
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    style_paragraph(p, after=after)


def add_bullets(doc, items, level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_run_font(p.add_run(item))
        style_paragraph(p, after=2)


def add_numbered(doc, items) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        set_run_font(p.add_run(f"{index}. {item}"))
        style_paragraph(p, after=3)


def add_code(doc, code: str) -> None:
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.right_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run_font(r, name="DejaVu Sans Mono", size=8.8, color="1F1F1F")
        set_cell_shading_paragraph(p, "F5F5F5")


def set_cell_shading_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    sizes = {1: 17, 2: 13.5, 3: 11.5}
    set_run_font(r, size=sizes[level], bold=True)
    style_paragraph(p, before=12 if level == 1 else 8, after=5, line=1.1, keep=True)


def add_table(doc, headers, rows, widths=None, font_size=9.2, landscape=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(str(h)), size=font_size, bold=True, color=WHITE)
        style_paragraph(p, after=0, line=1.05)
        if widths:
            cell.width = Cm(widths[i])
    for r_idx, row in enumerate(rows):
        added_row = table.add_row()
        set_row_cant_split(added_row)
        cells = added_row.cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) <= 4 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), size=font_size)
            style_paragraph(p, after=0, line=1.12)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=GRAY)


def make_equation_image(text: str, filename: str, width=1800, height=190, font_size=46) -> Path:
    path = WORK / filename
    image = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT_SANS, font_size)
    bbox = draw.textbbox((0, 0), text, font=font)
    x = (width - (bbox[2] - bbox[0])) // 2
    y = (height - (bbox[3] - bbox[1])) // 2 - bbox[1]
    draw.text((x, y), text, fill=f"#{BLACK}", font=font)
    image.save(path, dpi=(300, 300))
    return path


def make_flow_image() -> Path:
    path = WORK / "stage1_data_flow.png"
    width, height = 1800, 540
    image = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(FONT_SANS, 44)
    font = ImageFont.truetype(FONT_SANS, 32)
    small = ImageFont.truetype(FONT_SANS, 27)
    draw.text((70, 35), "Stage 1 训练与控制数据流", font=title_font, fill=f"#{BLACK}")
    labels = [
        ("28维观测", "q qd 跟踪误差\n末端 目标 碰撞"),
        ("PPO Actor", "两层64维 Tanh\n输出7维动作"),
        ("目标关节增量", "Δq = 0.02a\n累计为 q_target"),
        ("MuJoCo位置伺服", "kp kv 力矩限幅\n10个物理子步"),
        ("新状态与奖励", "距离 成功 碰撞\n返回下一步"),
    ]
    x0, y0, box_w, box_h, gap = 45, 185, 290, 215, 60
    for i, (name, detail) in enumerate(labels):
        x = x0 + i * (box_w + gap)
        draw.rounded_rectangle((x, y0, x + box_w, y0 + box_h), radius=22, fill="#F3F6FA", outline="#1F4E79", width=4)
        tb = draw.textbbox((0, 0), name, font=font)
        draw.text((x + (box_w - (tb[2] - tb[0])) / 2, y0 + 30), name, font=font, fill="#1F4E79")
        lines = detail.split("\n")
        for j, line in enumerate(lines):
            db = draw.textbbox((0, 0), line, font=small)
            draw.text((x + (box_w - (db[2] - db[0])) / 2, y0 + 100 + j * 42), line, font=small, fill="#333333")
        if i < len(labels) - 1:
            ax1 = x + box_w + 10
            ax2 = x + box_w + gap - 10
            ay = y0 + box_h // 2
            draw.line((ax1, ay, ax2, ay), fill="#1F4E79", width=6)
            draw.polygon([(ax2, ay), (ax2 - 18, ay - 12), (ax2 - 18, ay + 12)], fill="#1F4E79")
    image.save(path, dpi=(220, 220))
    return path


def make_results_chart() -> Path:
    path = WORK / "rm75_results_comparison.png"
    width, height = 1800, 880
    image = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    title = ImageFont.truetype(FONT_SANS, 43)
    font = ImageFont.truetype(FONT_SANS, 29)
    small = ImageFont.truetype(FONT_SANS, 25)
    draw.text((80, 35), "RM75确定性评估结果  种子42至141  共100回合", font=title, fill=f"#{BLACK}")
    categories = ["总体", "负Y外侧", "正Y外侧", "远X区域"]
    old = [69.0, 20.7, 85.3, 52.5]
    new = [63.0, 75.9, 20.6, 67.5]
    mix = [73.0, 31.0, 82.4, 55.0]
    colors = ["#4C78A8", "#E45756", "#72B7B2"]
    names = ["RM75 v3", "困难区微调最终模型", "75%旧权重加25%新权重"]
    left, top, bottom, chart_w = 150, 150, 740, 1520
    draw.line((left, top, left, bottom), fill="#555555", width=3)
    draw.line((left, bottom, left + chart_w, bottom), fill="#555555", width=3)
    for tick in range(0, 101, 20):
        y = bottom - int((bottom - top) * tick / 100)
        draw.line((left, y, left + chart_w, y), fill="#E5E5E5", width=2)
        draw.text((75, y - 16), str(tick), font=small, fill="#555555")
    group_w = chart_w / len(categories)
    bar_w = 78
    datasets = [old, new, mix]
    for i, cat in enumerate(categories):
        center = left + group_w * (i + 0.5)
        for j, data in enumerate(datasets):
            x1 = center + (j - 1) * (bar_w + 18) - bar_w / 2
            y1 = bottom - (bottom - top) * data[i] / 100
            draw.rectangle((x1, y1, x1 + bar_w, bottom), fill=colors[j])
            val = f"{data[i]:.1f}"
            vb = draw.textbbox((0, 0), val, font=small)
            draw.text((x1 + (bar_w - (vb[2] - vb[0])) / 2, y1 - 35), val, font=small, fill=f"#{BLACK}")
        cb = draw.textbbox((0, 0), cat, font=font)
        draw.text((center - (cb[2] - cb[0]) / 2, bottom + 25), cat, font=font, fill=f"#{BLACK}")
    lx = 230
    for i, name in enumerate(names):
        draw.rectangle((lx, 825, lx + 38, 857), fill=colors[i])
        draw.text((lx + 52, 820), name, font=small, fill=f"#{BLACK}")
        lx += 480
    image.save(path, dpi=(220, 220))
    return path


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size in (("Title", 28), ("Subtitle", 14), ("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        st = styles[style_name]
        st.font.name = "Noto Sans CJK SC"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(BLACK)
        if style_name.startswith("Heading"):
            st.font.bold = True

    title_p_pr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for section in doc.sections:
        footer = section.footer
        add_page_number(footer.paragraphs[0])


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(22)
    remove_paragraph_border(p)
    set_run_font(p.add_run("RM75机械臂强化学习随机目标到达"), size=28, bold=True)

    p2 = doc.add_paragraph(style="Subtitle")
    set_run_font(p2.add_run("Stage 1 汇报总结"), size=17, bold=True, color=NAVY)
    p2.paragraph_format.space_after = Pt(32)

    p3 = doc.add_paragraph()
    set_run_font(p3.add_run("UR5e 基线复现  RM75迁移  PPO训练与评估"), size=13, color="333333")
    p3.paragraph_format.space_after = Pt(68)

    meta = [
        "汇报范围  固定起点到随机目标点的仿真到达任务",
        "资料依据  UR5E_rl.odt  zhongyu_RL.pdf  当前项目代码与训练结果",
        "汇报日期  2026年9月30日",
    ]
    for line in meta:
        p = doc.add_paragraph()
        set_run_font(p.add_run(line), size=11, color=GRAY)
        p.paragraph_format.space_after = Pt(7)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    set_run_font(
        p.add_run("当前结论  已完成RM75仿真环境和PPO训练闭环  当前任务仍是末端点到点到达"),
        size=11,
        bold=True,
        color=NAVY,
    )
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    add_heading(doc, "汇报结构", 1)
    rows = [
        ("1", "Stage 1目标与完成情况"),
        ("2", "项目架构与代码调用关系"),
        ("3", "输入输出奖励与网络"),
        ("4", "从原始UR5代码到当前RM75系统的修改记录"),
        ("5", "奖励函数设计及其优化过程"),
        ("6", "平滑运动的实现与速度参数说明"),
        ("7", "PPO与TD3的选择"),
        ("8", "训练结果与问题分析"),
        ("9", "结果衡量方法"),
        ("10", "Stage 1局限与下一步仿真优化"),
        ("附录", "代码位置 运行命令和汇报问答"),
    ]
    add_table(doc, ["章节", "内容"], rows, widths=[2.3, 13.4], font_size=10)
    add_text(
        doc,
        "范围说明：本报告只覆盖 zhongyu_RL.pdf 中的 Stage 1。真机测试暂不实施；障碍物点云、ESDF、多边形表示和语义接触属于后续阶段。",
        bold_lead="范围说明：",
    )
    doc.add_page_break()


def build_document() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    flow_img = make_flow_image()
    result_img = make_results_chart()
    action_eq = make_equation_image(
        "a(t) in [-1, 1]^7    delta_q(t) = 0.02 a(t)    q_target(t+1) = clip(q_target(t) + delta_q(t), q_min, q_max)",
        "action_equation.png",
        font_size=39,
    )
    reward_eq = make_equation_image(
        "r(t) = 50[d(t-1) - d(t)] - 0.03d(t) - 0.01 + r_milestone - 75I_collision + 100I_success - 10I_timeout",
        "reward_equation.png",
        height=210,
        font_size=35,
    )

    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "1 Stage 1目标与完成情况", 1)
    add_text(
        doc,
        "Stage 1 的目标是让机械臂从固定笛卡尔起点出发，在二维桌面工作区域内接收随机目标位置，并由强化学习策略生成连续的七关节控制序列，使末端执行器进入目标点附近。当前实现完成了 UR5e 基线复现、RM75 MuJoCo模型迁移、Gymnasium环境、PPO训练、模型评估和断点续训。",
    )
    rows = [
        ("本地部署与基线复现", "完成", "原仓库环境、PPO主流程和可视化逻辑已梳理并修复运行问题"),
        ("UR5e迁移到RM75", "完成", "七轴URDF、STL、MJCF、执行器、TCP临时site和场景已建立"),
        ("固定起点到随机目标", "完成", "RM75-v0提供28维观测、7维动作和随机目标采样"),
        ("训练与可视化评估", "完成", "支持W&B、TensorBoard、批量确定性评估和窗口演示"),
        ("真机测试", "本阶段暂不实施", "仿真增益、TCP和关节方向不能直接当作实机参数"),
    ]
    add_table(doc, ["工作项", "状态", "说明"], rows, widths=[4.0, 2.8, 9.0], font_size=9.4)
    add_text(
        doc,
        "本阶段完成的是点到点 reaching，不是随机物体抓取。环境没有夹爪闭合动作、物体状态、抓取保持或抬升成功判据；碰撞检测用于结束危险回合，也不等同于 Stage 2 的障碍物感知规划。",
        bold_lead="能力边界：",
    )

    add_heading(doc, "2 项目架构与代码调用关系", 1)
    add_text(
        doc,
        "项目按模型资产、仿真环境、强化学习算法、评估脚本和实验结果分层。PPO不直接修改MuJoCo状态，只通过Gymnasium的reset和step接口交换数据。环境负责把归一化动作转换为目标关节角、推进物理仿真并计算奖励。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(flow_img), width=Cm(16.2))
    style_paragraph(p, after=2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run("图1  Stage 1训练与控制数据流"), size=9, color=GRAY)
    style_paragraph(cap, after=8)

    rows = [
        ("assets/rm75/RM75.urdf", "厂家ROS模型的项目副本，网格路径改为相对路径"),
        ("assets/rm75/rm75_imported.xml", "URDF自动转换结果，用于比较，不直接作为训练模型"),
        ("assets/rm75/rm75.xml", "审核后的RM75动力学、执行器、限位、home姿态和临时TCP"),
        ("assets/rm75/scene.xml", "地面、灯光和相机视觉设置"),
        ("env_rm75.py", "任务定义、初始化、IK、控制、碰撞、奖励和渲染"),
        ("Algorithms/ppo_continuous_action.py", "Actor Critic、rollout、GAE、PPO更新、日志与断点"),
        ("Scripts/visualize_agent.py", "加载训练权重，窗口演示或无界面批量评估"),
        ("tests", "模型、环境、奖励、碰撞、PPO和checkpoint契约测试"),
    ]
    add_table(doc, ["文件", "职责"], rows, widths=[6.3, 9.5], font_size=9.2)
    add_heading(doc, "2.1 运行时调用顺序", 2)
    add_numbered(doc, [
        "训练脚本解析参数并注册 RM75-v0。",
        "gym.make 创建 RM75P2PEnv，环境读取 scene.xml 和 rm75.xml，并编译 MjModel。",
        "reset 清空 MjData，随机化关节初值和桌面摩擦，通过IK把末端放到固定起点附近。",
        "Actor读取28维观测并生成7维归一化动作。",
        "环境把动作乘以0.02 rad后累计到q_target，位置执行器在10个物理子步内跟踪目标。",
        "环境读取新状态、距离和碰撞，返回奖励和终止标志。",
        "PPO收集2048步rollout，计算GAE和return，再进行小批量更新。",
    ])

    add_heading(doc, "3 输入输出奖励与网络", 1)
    add_heading(doc, "3.1 输入  28维观测", 2)
    obs_rows = [
        ("obs[0:7]", "7", "当前关节角 q", "描述机械臂构型"),
        ("obs[7:14]", "7", "当前关节速度 qd", "让策略感知运动方向和动态状态"),
        ("obs[14:21]", "7", "位置伺服跟踪误差 q_target - q", "暴露累计控制目标，消除隐藏状态"),
        ("obs[21:24]", "3", "末端世界坐标", "给出当前位置"),
        ("obs[24:27]", "3", "目标世界坐标", "告诉策略需要到达的位置"),
        ("obs[27]", "1", "碰撞标志", "表示当前接触状态"),
    ]
    add_table(doc, ["切片", "维数", "含义", "作用"], obs_rows, widths=[3.0, 1.5, 5.0, 6.3], font_size=8.8)

    add_heading(doc, "3.2 输出  7维关节位置参考增量", 2)
    add_text(
        doc,
        "策略输出不是力矩，也不是关节速度，而是无量纲动作。环境把它转换为每个环境步的目标关节角增量，并累加到位置伺服参考值。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(action_eq), width=Cm(16.0))
    style_paragraph(p, after=4)
    add_text(
        doc,
        "由于动作空间为 [-1,1]，单步目标变化不超过 ±0.02 rad。这个限制既统一了策略概率空间，也避免一次动作给出过大的关节跳变。",
    )

    add_heading(doc, "3.3 PPO Actor Critic网络", 2)
    rows = [
        ("Actor均值网络", "28 → 64 → 64 → 7", "两层Tanh，输出高斯分布均值"),
        ("策略标准差", "7个可学习log std", "受训练进度上限约束，训练动作从高斯分布采样"),
        ("动作有界化", "tanh压缩", "动作严格落入(-1,1)，log probability包含变量变换修正"),
        ("Critic", "28 → 64 → 64 → 1", "估计状态价值V(s)，用于优势函数和PPO更新"),
        ("初始化", "正交初始化", "改善网络初始尺度和训练稳定性"),
    ]
    add_table(doc, ["模块", "结构", "说明"], rows, widths=[3.7, 4.2, 8.0], font_size=9.1)

    add_heading(doc, "4 从原始UR5代码到当前RM75系统的修改记录", 1)
    add_text(
        doc,
        "下表按实际迭代顺序汇总主要修改。每一项都给出文件或函数位置、修改内容、原因和结果。这里的“每一次修改”按可验证的功能批次归纳，而不是逐行罗列格式变化。",
    )
    change_rows = [
        ("1", "env.py 主入口", "交互模式创建环境时固定 render_mode=None，由 interactive_viewer 单独启动窗口", "避免human渲染和交互查看器重复launch造成段错误", "交互模式可与target-mode同时使用"),
        ("2", "PPO入口与sys.path", "删除不存在的UR5包副作用导入，直接导入环境类并注册ID", "修复ModuleNotFoundError", "脚本可从项目根目录直接运行"),
        ("3", "env.py action_space与step", "动作改为[-1,1]归一化，再乘0.02 rad", "策略分布单位与物理单位原来混在一起", "动作概率计算和实际执行一致"),
        ("4", "PPO Agent", "高斯采样后用tanh压缩，并修正log probability；删除ClipAction", "事后裁剪会使环境动作与PPO记录动作不一致", "保证PPO概率比计算正确"),
        ("5", "env.py reset", "模型只编译一次，回合重置改用mj_resetData", "原代码每回合重建XML模型，速度慢且随机性难管理", "训练吞吐提高，模型ID保持不变"),
        ("6", "env.py reset与IK", "先定义start_pos，再对带噪q0执行IK", "IK必须知道目标起点；噪声提供初态随机化", "末端每回合回到起点附近但关节构型有扰动"),
        ("7", "PPO默认参数", "学习率降至1e-4，rollout增至2048，更新轮数降至5，target KL设0.02", "原更新过强，曲线波动大", "限制单批策略变化"),
        ("8", "PPO日志", "增加return length success collision timeout final distance saturation和区域成功率", "只看loss无法判断任务是否完成", "训练可诊断且可比较"),
        ("9", "Scripts/visualize_agent.py", "修复导入；加入自动选模型、确定性动作、headless和汇总统计", "原脚本不能加载本项目环境且只适合观看", "可复现评估随机目标性能"),
        ("10", "tools/convert_rm75_urdf.py", "复制厂家URDF和STL，替换package路径并导出基础MJCF", "MuJoCo不能直接解析ROS package路径", "转换流程可重复"),
        ("11", "rm75_imported.xml与rm75.xml", "自动导入文件和人工审核训练文件分离", "再次转换不应覆盖执行器和人工参数", "模型更新可追踪"),
        ("12", "assets/rm75/rm75.xml", "加入7个位置执行器、关节限位、力矩限幅、home关键帧和ee_site", "URDF导入结果没有可训练控制接口", "形成7轴可控模型"),
        ("13", "assets/rm75/scene.xml", "加入地面、灯光、视觉设置；训练桌子由环境动态加入", "静态场景和任务随机量需要分离", "摩擦可以由reset种子控制"),
        ("14", "env_rm75.py", "建立独立RM75P2PEnv，不覆盖UR5环境", "UR5为6轴而RM75为7轴，动力学和限位不同", "两套环境可以并存"),
        ("15", "_build_scene_and_model", "编译后核验joint_1至joint_7和对应actuator顺序", "顺序错位会把策略动作发送到错误关节", "运行时快速失败而不是静默训练错误"),
        ("16", "_randomize_table_friction", "使用self.rng在reset时更新geom_friction", "原全局随机数不受seed控制；重编译模型代价高", "同seed可复现且无需重编译"),
        ("17", "_check_collision", "加入桌面、地面和非相邻连杆自碰撞；相邻连接接触不算碰撞", "只检查桌面会漏掉机械臂自碰撞", "碰撞终止含义更完整"),
        ("18", "Gymnasium注册", "保留UR5-v0并新增RM75-v0，重复导入时不重复注册", "迁移不应破坏基线", "gym.make可创建两种机械臂"),
        ("19", "PPO网络维度", "从环境空间动态获得输入输出维度", "不能硬编码UR5的19维和6维", "同一PPO代码支持28维7动作RM75"),
        ("20", "观测V2", "加入q_target-q，21维扩展为28维", "累计位置目标原来是策略看不到的隐藏状态", "状态更接近马尔可夫性"),
        ("21", "奖励V2", "加入超时惩罚并加强距离约束", "长回合无成功时缺乏明确失败信号", "试验后发现约束过强，继续进入V3"),
        ("22", "奖励V3", "timeout改10，dist_weight改0.03，加入20cm和10cm一次性里程碑", "V2成功样本极少，终端成功奖励太稀疏", "策略获得分阶段接近信号"),
        ("23", "探索V3", "标准差前25%保持0.30，再降至0.12；entropy系数为0.0001", "过早降低探索会很难首次采到成功轨迹", "早期充分探索，后期动作更确定"),
        ("24", "目标采样", "新增uniform和50/25/25 hard_mixture，并记录target_region", "诊断出负Y和远X区域成功率低", "可定向微调，但当前结果出现区域遗忘"),
        ("25", "checkpoint", "区分init-from和resume-from；原子保存Agent Adam步数参数和随机状态", "纯权重不能真正续训，Ctrl+C后也可能丢失进度", "支持微调和中断恢复"),
        ("26", "测试体系", "建立26项模型环境奖励PPO和checkpoint单元测试", "迁移涉及维度顺序碰撞和随机性，人工观看不足以验证", "当前26项测试全部通过"),
    ]
    add_table(
        doc,
        ["序号", "修改位置", "修改内容", "原因", "影响"],
        change_rows,
        widths=[1.0, 3.1, 4.4, 4.0, 3.5],
        font_size=7.9,
    )

    add_heading(doc, "5 奖励函数设计及其优化过程", 1)
    add_text(
        doc,
        "当前奖励集中定义在 env_rm75.py 的 reward_cfg，并由 _compute_reward 统一计算。step先计算末端到目标的欧氏距离、碰撞、成功和超时，再把各分量写入 info.reward_components，便于训练诊断。",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(reward_eq), width=Cm(16.0))
    style_paragraph(p, after=2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run("式中 d 表示末端与目标的欧氏距离  I表示条件成立时取1"), size=8.8, color=GRAY)
    style_paragraph(cap, after=7)

    reward_rows = [
        ("进度奖励", "50(dₜ₋₁-dₜ)", "只奖励距离净减少；远离目标时为负。前进再退回的两步收益基本抵消，降低来回摆动获利的可能。"),
        ("距离惩罚", "-0.03dₜ", "让未产生明显进度时仍偏向靠近目标。权重保持较小，避免长回合累计距离项压过成功奖励。"),
        ("时间惩罚", "每步-0.01", "同样成功时偏向使用更少步骤，但强度较小，不会主导策略。"),
        ("20cm里程碑", "首次进入时+5", "给稀疏成功奖励增加中间学习信号；每回合只触发一次。"),
        ("10cm里程碑", "首次进入时+15", "越接近目标给出更强信号；来回跨越阈值不能重复领取。"),
        ("碰撞惩罚", "碰撞时-75", "碰撞同时终止回合，明确区分不安全轨迹与普通超时。"),
        ("成功奖励", "d<0.05m时+100", "把进入5cm目标球定义为本阶段任务成功，并立即终止回合。"),
        ("超时惩罚", "失败到1024步时-10", "防止长期停滞；只在未成功且未碰撞时施加，避免重复惩罚。"),
        ("姿态奖励", "当前权重为0", "代码接口保留，但Stage 1只训练位置到达，不声称已学习末端姿态。"),
    ]
    add_table(doc, ["分量", "当前形式", "设计理由"], reward_rows, widths=[3.0, 3.3, 9.7], font_size=8.9)

    add_heading(doc, "5.1 为什么改成势函数式进度奖励", 2)
    add_text(
        doc,
        "原实现使用 10(d_prev-d)/max(d,10⁻⁶)。当d很小时，分母会把很小的距离变化放大，策略可能在目标附近通过数值波动获得异常奖励。当前实现使用50(d_prev-d)，其累计值主要由初始距离和最终距离决定。向前0.02m获得约+1，随后退回0.02m得到约-1，因此往返本身没有持续收益。由于PPO使用γ=0.99，严格数学抵消会有轻微折扣误差，但比按当前距离相除更稳定。",
    )
    add_heading(doc, "5.2 奖励演化", 2)
    reward_history = [
        ("原始UR5", "相对进度除以当前距离，距离权重0.5，碰撞10", "近目标奖励可能放大，距离累计项较强"),
        ("UR5稳定修复", "线性进度50，距离0.02，碰撞70", "先减少摆动获利和碰撞行为"),
        ("RM75 V2", "timeout25，距离0.1，无里程碑", "失败信号明确，但稀疏成功和强负项导致成功很少"),
        ("RM75 V3当前", "timeout10，距离0.03，里程碑5和15", "保留目标方向，同时提高采到成功轨迹后的可学习性"),
    ]
    add_table(doc, ["版本", "关键设置", "判断"], reward_history, widths=[3.3, 6.3, 6.4], font_size=9.0)
    add_heading(doc, "5.3 下一步怎样优化奖励", 2)
    add_bullets(doc, [
        "先不要继续改成功、碰撞和距离权重。最新实验的主要问题是目标采样失衡，而不是奖励完全无效。",
        "加入轨迹平滑指标后，可小权重增加动作变化惩罚，例如 -λΔa‖aₜ-aₜ₋₁‖²；必须同时监控成功率，避免策略因害怕动作而停住。",
        "如果目标从位置扩展为位姿，再启用姿态误差奖励；在Stage 1位置任务中保持orient_weight=0。",
        "每次只改一组因素，并用固定验证种子比较总体成功率、最弱区域成功率、碰撞率、步数和轨迹平滑度。",
    ])

    add_heading(doc, "6 平滑运动的实现与速度参数说明", 1)
    add_text(
        doc,
        "当前系统没有显式的轨迹规划器，也没有给PPO设置一个名为joint_speed的命令参数。运动看起来较平滑，主要来自目标关节增量限制、位置伺服器动力学、物理子步和探索噪声衰减。它是受约束的闭环响应，不是经过速度、加速度和加加速度规划后得到的严格平滑轨迹。",
    )
    smooth_rows = [
        ("动作增量限制", "env_rm75.py  action_scale=0.02", "每20ms只允许目标关节角最多变化0.02rad，避免参考值瞬间大跳"),
        ("累计位置目标", "step中的q_target", "策略给增量，执行器持续追踪累计目标；未完成的跟踪不会在下一步丢失"),
        ("位置伺服阻尼", "rm75.xml中的position actuator", "1至2轴kp200 kv20，3至4轴kp120 kv12，5至7轴kp40 kv4；kv对速度产生阻尼"),
        ("力矩限幅", "forcerange", "各轴限制为60 60 30 30 10 10 10Nm，避免仿真执行器无限施力"),
        ("物理子步", "timestep=0.002s  substeps=10", "每个策略动作内部推进10次动力学积分，环境控制频率为50Hz"),
        ("速度观测", "obs[7:14]=qvel", "速度存在于观测，策略可以学会利用，但奖励没有直接惩罚高速"),
        ("探索退火", "action std schedule", "训练后期随机动作减小，轨迹通常更稳定；评估默认用确定性均值动作"),
    ]
    add_table(doc, ["机制", "代码位置", "对平滑性的作用"], smooth_rows, widths=[3.3, 5.2, 7.5], font_size=8.9)

    add_heading(doc, "6.1 速度在哪里", 2)
    add_text(
        doc,
        "MuJoCo中的实际关节速度是 data.qvel，由动力学积分产生，并作为观测输入策略。当前动作并不是速度命令。按参考值计算，单步最大增量0.02rad、环境步长0.02s，对应目标参考值最快可连续改变约1rad/s；这个1rad/s只是参考变化上限，不是机械臂真实速度限值。实际速度还受惯量、kp、kv、力矩限幅和接触影响。",
    )
    add_text(
        doc,
        "因此，不能仅凭action_scale就声称满足实机速度安全要求。若要形成可验证的平滑轨迹，需要加入关节速度、加速度和jerk统计，并在环境或低层控制器中显式限制。",
        bold_lead="关键结论：",
    )
    add_heading(doc, "6.2 下一步平滑性改进", 2)
    add_numbered(doc, [
        "在info和评估脚本中记录每回合max|qd|、RMS速度、RMS加速度和RMS jerk。",
        "根据厂家仿真目标值加入每轴速度和加速度限制；真机参数暂不使用。",
        "优先增加动作变化惩罚或低通滤波，再观察是否降低成功率。",
        "如需严格时间最优且平滑，可让策略输出末端或关节目标，再由独立轨迹生成器做速度规划。",
    ])

    add_heading(doc, "7 PPO与TD3的选择", 1)
    comparison = [
        ("策略类型", "随机策略  on policy", "确定性策略  off policy"),
        ("数据使用", "采样后更新若干轮，之后丢弃", "经验回放可重复利用，样本效率通常更高"),
        ("稳定机制", "概率比裁剪、KL阈值、优势标准化", "双Critic、延迟Actor更新、目标策略平滑"),
        ("探索", "策略分布标准差直接控制", "确定性动作外加探索噪声"),
        ("当前工程成熟度", "已完成RM75接入、日志、测试、评估和checkpoint", "仓库有代码，但尚未完成同等RM75验证"),
        ("Stage 1建议", "作为当前主算法", "后续作为对照实验，不直接替换"),
    ]
    add_table(doc, ["比较项", "PPO", "TD3"], comparison, widths=[3.0, 6.5, 6.5], font_size=9.0)
    add_text(
        doc,
        "选择PPO的直接原因是当前任务在仿真中可大量采样，PPO更新幅度容易限制，现有训练和诊断链路已经完整。TD3的经验回放更适合昂贵数据场景，但需要重新验证回放分布、探索噪声、Q值过估计和RM75动作尺度。Stage 1应先用PPO建立可靠基线，再用相同目标种子和相同评价指标做TD3对照。",
    )

    add_heading(doc, "8 训练结果与问题分析", 1)
    run_rows = [
        ("RM75初始100万步", "999424", "末段滚动成功14%  最高32%", "碰撞和超时较多，基础策略尚不稳定"),
        ("RM75 V2", "约108万", "末段滚动成功0%  最高约1%", "超时和距离负项过强，成功奖励过于稀疏"),
        ("RM75 V3", "1998848", "末段滚动成功79%  碰撞0%", "奖励和探索修正后学会任务"),
        ("困难区微调", "499712", "末段滚动成功64%  碰撞0%", "负Y改善，但正Y发生遗忘"),
    ]
    add_table(doc, ["实验", "实际环境步", "训练曲线", "解释"], run_rows, widths=[3.6, 2.7, 4.4, 5.3], font_size=8.9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(result_img), width=Cm(16.2))
    style_paragraph(p, after=2)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run("图2  相同100个随机目标下的区域成功率对比"), size=9, color=GRAY)
    style_paragraph(cap, after=8)

    eval_rows = [
        ("RM75 v3", "69%", "0%", "0.1232m", "463.7"),
        ("困难区微调最终模型", "63%", "0%", "0.1134m", "503.7"),
        ("75%旧权重加25%新权重", "73%", "0%", "0.1115m", "445.3"),
    ]
    add_table(doc, ["模型", "成功率", "碰撞率", "平均最终距离", "平均步数"], eval_rows, widths=[5.0, 2.2, 2.2, 3.3, 2.8], font_size=9.0)
    add_text(
        doc,
        "困难区微调没有整体提升，而是把能力从正Y区域迁移到负Y区域。负Y外侧成功率从20.7%升至75.9%，正Y外侧却从85.3%降至20.6%。另一组独立100种子也复现了这一趋势，因此主要问题是50/25/25采样分布导致的区域遗忘。",
        bold_lead="主要结论：",
    )
    add_text(
        doc,
        "旧模型75%与新模型25%的权重插值在两组100种子测试中分别取得73%和80%，高于对应旧模型的69%和78%。这是有价值的候选，但尚未保存为正式模型，也需要用至少500个全新种子复核。",
    )

    add_heading(doc, "9 结果衡量方法", 1)
    metric_rows = [
        ("成功率", "成功回合数除以总回合数", "核心任务指标；训练滚动值与独立评估值必须区分"),
        ("分区成功率", "按负Y 正Y 远X等区域统计", "发现总体均值掩盖的空间盲区"),
        ("碰撞率", "碰撞终止回合比例", "评价安全性；当前确定性评估为0"),
        ("超时率", "1024步未成功且未碰撞的比例", "判断是否停滞或无法到达"),
        ("最终距离", "终止时末端到目标距离", "失败策略是否仍在接近目标"),
        ("回合步数", "成功或失败前使用的环境步数", "成功率相同时衡量效率"),
        ("动作饱和率", "|a|>0.98的动作比例", "判断策略是否长期顶住动作边界"),
        ("平滑度", "RMS速度 加速度 jerk和动作差", "当前尚未加入正式日志，是下一步补充项"),
        ("PPO诊断", "KL clipfrac entropy value loss explained variance", "解释更新是否过强和Critic预测质量，不能代替成功率"),
    ]
    add_table(doc, ["指标", "定义", "用途"], metric_rows, widths=[3.1, 5.0, 7.9], font_size=8.9)
    add_heading(doc, "9.1 规范评估流程", 2)
    add_numbered(doc, [
        "训练曲线用于观察学习过程，不直接当作最终性能。",
        "评估默认使用tanh(actor_mean)确定性动作，排除训练探索噪声。",
        "使用固定且未参与调参的随机种子，至少评估500个目标。",
        "同时报告总体和区域指标，不只报告最好一次回合。",
        "保存模型、环境参数、代码版本和种子，使结果可以复现。",
    ])

    add_heading(doc, "10 Stage 1局限与下一步仿真优化", 1)
    limitation_rows = [
        ("临时TCP", "ee_site位于link_7原点，没有真实夹爪变换", "仿真点到达有效，但不能代表真实工具中心"),
        ("目标只有位置", "目标z固定，未训练姿态", "不能完成需要末端朝向的抓取预备位姿"),
        ("没有夹爪和物体", "没有闭合、接触保持、抬升判据", "不能声称完成随机点抓取"),
        ("平滑性未显式优化", "没有速度加速度jerk限制或惩罚", "视觉上平滑不等于轨迹指标满足要求"),
        ("碰撞信息简化", "观测只有一个碰撞布尔量", "策略不知道障碍方向和距离；这属于后续阶段"),
        ("随机目标采样偏置", "hard_mixture过度强化负Y", "造成正Y能力遗忘"),
        ("仿真参数来源", "质量惯量采用当前URDF，执行器增益为仿真值", "不能直接映射到真机"),
    ]
    add_table(doc, ["局限", "当前状态", "影响"], limitation_rows, widths=[3.2, 6.0, 6.8], font_size=8.9)

    add_heading(doc, "10.1 推荐下一步", 2)
    next_rows = [
        ("1", "改为均衡目标采样", "60%全区域 15%负Y 15%正Y 10%远X", "维持困难区训练，同时避免左右区域遗忘"),
        ("2", "增加周期验证和best model", "每5万步在固定200个均匀目标上评估", "最终权重不一定是训练过程中最好权重"),
        ("3", "验证插值候选", "至少500个全新种子比较旧 新 插值模型", "确认73%和80%的提升不是偶然"),
        ("4", "小学习率均衡微调", "约30万步  学习率1e-5  std从0.08降至0.06", "恢复正Y时减少对现有能力的破坏"),
        ("5", "加入平滑度日志", "速度 加速度 jerk 动作差", "把视觉判断变成可比较的量化结果"),
        ("6", "保持奖励暂时不变", "先隔离采样因素", "避免同时改采样和奖励后无法解释结果"),
    ]
    add_table(doc, ["优先级", "工作", "建议", "原因"], next_rows, widths=[1.5, 3.5, 5.4, 5.6], font_size=8.9)

    add_heading(doc, "附录A 关键代码位置", 1)
    code_rows = [
        ("动作和观测空间", "env_rm75.py  RM75P2PEnv.__init__"),
        ("模型编译与顺序检查", "env_rm75.py  _build_scene_and_model"),
        ("初始化和IK", "env_rm75.py  reset  _solve_ik"),
        ("动作到关节目标", "env_rm75.py  step  _pd_to_ctrl  _step_sim"),
        ("奖励", "env_rm75.py  reward_cfg  _compute_reward  _compute_distance_milestone_reward"),
        ("观测", "env_rm75.py  _get_obs"),
        ("碰撞", "env_rm75.py  _find_robot_geom_ids  _contact_pair_is_collision  _check_collision"),
        ("网络", "Algorithms/ppo_continuous_action.py  Agent"),
        ("PPO更新", "Algorithms/ppo_continuous_action.py  rollout  GAE  minibatch update"),
        ("断点", "save_training_checkpoint  load_training_checkpoint"),
        ("评估", "Scripts/visualize_agent.py"),
        ("模型参数", "assets/rm75/rm75.xml"),
    ]
    add_table(doc, ["主题", "代码位置"], code_rows, widths=[5.4, 10.6], font_size=9.1)

    add_heading(doc, "附录B 可复制运行命令", 1)
    add_heading(doc, "B.1 完整自动测试", 2)
    add_code(doc, "RM75_DEBUG=0 .venv/bin/python -m unittest discover -s tests -v")
    add_heading(doc, "B.2 从头训练RM75 V3", 2)
    add_code(doc, """RM75_DEBUG=0 .venv/bin/python Algorithms/ppo_continuous_action.py \\
  --env-id RM75-v0 \\
  --total-timesteps 2000000 \\
  --learning-rate 1e-4 \\
  --track""")
    add_heading(doc, "B.3 批量确定性评估", 2)
    add_code(doc, """.venv/bin/python Scripts/visualize_agent.py \\
  --model-path runs/RM75-v3/ppo_continuous_action.cleanrl_model \\
  --headless \\
  --episodes 500 \\
  --seed 20000 \\
  --progress-every 0""")
    add_heading(doc, "B.4 窗口演示", 2)
    add_code(doc, """.venv/bin/python Scripts/visualize_agent.py \\
  --model-path runs/RM75-v3/ppo_continuous_action.cleanrl_model \\
  --episodes 5""")
    add_heading(doc, "B.5 中断恢复", 2)
    add_code(doc, """.venv/bin/python Algorithms/ppo_continuous_action.py \\
  --resume-from runs/RUN_NAME/checkpoints/checkpoint_latest.pt""")
    add_text(
        doc,
        "训练中第一次按Ctrl+C会完成当前PPO迭代后保存checkpoint。不要使用Ctrl+Z，Ctrl+Z只会挂起进程。已完整结束的模型开启新一轮优化时使用--init-from，而不是--resume-from。",
    )

    add_heading(doc, "附录C 老师可能追问的简短回答", 1)
    qa_rows = [
        ("输入是什么", "28维状态，包括关节角、关节速度、伺服跟踪误差、末端位置、目标位置和碰撞标志。"),
        ("输出是什么", "7维归一化动作，经0.02缩放后作为每步目标关节角增量，不是力矩或速度命令。"),
        ("奖励为什么这样设计", "终端成功奖励定义任务，进度和里程碑解决稀疏反馈，时间项鼓励效率，碰撞和超时区分失败类型。"),
        ("为什么选择PPO", "仿真可大量采样，PPO裁剪更新稳定，当前工程链路已经完整；TD3留作同指标对照。"),
        ("速度参数在哪里", "实际速度在MuJoCo的qvel中；动作是位置参考增量。0.02rad和50Hz给出约1rad/s参考变化上限，但不是实际速度限制。"),
        ("运动为什么比较平滑", "小增量、位置伺服kp kv、力矩限幅、10个物理子步和后期低探索共同作用；当前没有显式jerk优化。"),
        ("成功率为什么会下降", "困难采样让策略重点适应负Y，正Y样本不足，引起区域遗忘；不是简单的训练步数不够。"),
        ("现在能不能抓取", "不能。当前只验证末端点到达，没有夹爪、物体状态和抓取成功判据。"),
        ("真机是否完成", "本阶段暂不进行真机测试，仿真TCP、增益和方向仍需后续标定。"),
    ]
    add_table(doc, ["问题", "回答"], qa_rows, widths=[4.1, 11.9], font_size=9.0)

    add_heading(doc, "参考资料", 1)
    refs = [
        "zhongyu_RL.pdf  Stage 1任务要求与汇报问题",
        "UR5E_rl.odt  原始学习总结",
        "UR5-Reinforcement-Learning当前本地代码  env.py  env_rm75.py  PPO与评估脚本",
        "docs/superpowers/plans中的迁移 稳定性 奖励和断点实现记录",
        "runs目录中的TensorBoard事件 模型权重和checkpoint",
    ]
    add_numbered(doc, refs)

    for section in doc.sections:
        footer = section.footer
        if not footer.paragraphs[0].text:
            add_page_number(footer.paragraphs[0])

    doc.core_properties.title = "RM75机械臂强化学习随机目标到达 Stage 1 汇报总结"
    doc.core_properties.subject = "UR5e基线复现 RM75迁移 PPO训练与评估"
    doc.core_properties.author = "ZhongYu"
    doc.core_properties.keywords = "RM75, UR5e, PPO, MuJoCo, Reinforcement Learning"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
